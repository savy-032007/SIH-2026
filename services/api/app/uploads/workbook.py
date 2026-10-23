"""Building upload workbooks from the column contract (SPEC 10.1.3).

Used by the synthetic generator in Phase 1 and by the template-download endpoint
in Phase 2, so the fixtures and the real template are the same artefact.
"""

from __future__ import annotations

from collections.abc import Iterable, Mapping
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.uploads.template import (
    COLUMNS,
    DATA_SHEET_NAME,
    FIRST_DATA_ROW,
    HEADER_KEY_ROW,
    HEADER_LABEL_ROW,
    LISTS_SHEET_NAME,
    english_labels,
)

#: Leading characters Excel and LibreOffice treat as the start of a formula.
#: Values beginning with one of these are prefixed on write and stripped on read
#: (SPEC section 8, CSV injection).
FORMULA_TRIGGERS = ("=", "+", "-", "@")


def neutralize_formula(value: Any) -> Any:
    """Defang a cell value that a spreadsheet would evaluate as a formula.

    Applied on the way *out* as well as in: a generated fixture containing a live
    formula would be a self-inflicted version of the attack the parser guards
    against.
    """
    if isinstance(value, str) and value.startswith(FORMULA_TRIGGERS):
        return "'" + value
    return value


def build_workbook(
    rows: Iterable[Mapping[str, Any]],
    *,
    labels: Mapping[str, str] | None = None,
    include_lists_sheet: bool = True,
) -> Workbook:
    """Build a workbook with the two header rows and one row per record.

    `labels` supplies translated header labels for row 1; row 2 always carries the
    English machine keys the parser reads.
    """
    resolved_labels = dict(english_labels())
    if labels:
        resolved_labels.update(labels)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET_NAME

    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.cell(row=HEADER_LABEL_ROW, column=index, value=resolved_labels.get(column.key, column.key))
        sheet.cell(row=HEADER_KEY_ROW, column=index, value=column.key)
        sheet.column_dimensions[letter].width = column.width

    sheet.freeze_panes = sheet.cell(row=FIRST_DATA_ROW, column=1)

    row_number = FIRST_DATA_ROW
    for record in rows:
        for index, column in enumerate(COLUMNS, start=1):
            sheet.cell(
                row=row_number,
                column=index,
                value=neutralize_formula(record.get(column.key)),
            )
        row_number += 1

    if include_lists_sheet:
        _add_lists_sheet(workbook, sheet, last_data_row=max(row_number + 500, FIRST_DATA_ROW + 500))

    return workbook


def _add_lists_sheet(workbook: Workbook, data_sheet: Worksheet, *, last_data_row: int) -> None:
    """Add the hidden dropdown-source sheet and wire validations to it.

    Dropdowns are sourced from a sheet rather than inline lists because Excel caps
    inline validation formulas at 255 characters, which the nationality and board
    lists exceed.
    """
    lists = workbook.create_sheet(LISTS_SHEET_NAME)

    column_index = 1
    for column in COLUMNS:
        if column.kind != "enum" or not column.choices:
            continue

        letter = get_column_letter(column_index)
        lists.cell(row=1, column=column_index, value=column.key)
        for offset, choice in enumerate(column.choices, start=2):
            lists.cell(row=offset, column=column_index, value=choice)

        source = f"'{LISTS_SHEET_NAME}'!${letter}$2:${letter}${len(column.choices) + 1}"
        validation = DataValidation(type="list", formula1=source, allow_blank=not column.required)
        validation.error = f"Choose a value from the {column.key} list"
        validation.errorTitle = "Invalid value"
        data_sheet.add_data_validation(validation)

        target_letter = get_column_letter(COLUMNS.index(column) + 1)
        validation.add(f"{target_letter}{FIRST_DATA_ROW}:{target_letter}{last_data_row}")

        column_index += 1

    lists.sheet_state = "hidden"
